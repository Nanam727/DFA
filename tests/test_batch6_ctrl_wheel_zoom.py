"""批次6：Ctrl+鼠标滚轮调整预览缩放比例。

覆盖：
- Ctrl+滚轮上/下 → 缩放 +10%/-10%（Windows MouseWheel 与 Linux Button-4/5 两种事件形态）
- 缩放钳制在 50%–200%
- 普通滚轮在按住 Ctrl 时不再翻页/滚动（让位给缩放快捷键）
- 画布、标注栏、两个页面文本控件上都绑定了 Control 绑定
- 撑满宽模式下缩放真实生效（页面变宽、出现横向滚动条）
- 端到端：event_generate 触发绑定后缩放确实变化
"""

from __future__ import annotations

import sys
import tempfile
import types
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document

from src.ui import FormatApp

CTRL = 0x0004

PASS = 0
FAIL = 0


def _check(name: str, cond: bool, detail: str = "") -> None:
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  PASS  {name}")
    else:
        FAIL += 1
        print(f"  FAIL  {name}  {detail}")


def _wheel(delta: int = 120, num: int = 0, state: int = CTRL):
    return types.SimpleNamespace(delta=delta, num=num, state=state)


def test_ctrl_wheel_zoom_steps() -> None:
    print("[1] Ctrl+滚轮步进缩放")
    app = FormatApp()
    app.withdraw()
    try:
        app._set_preview_zoom(100)
        app._on_ctrl_mousewheel_zoom(_wheel(delta=120))
        _check("滚轮向上放大一档", app.preview_zoom_percent == 110,
               f"zoom={app.preview_zoom_percent}")
        app._on_ctrl_mousewheel_zoom(_wheel(delta=120))
        _check("再向上一档到 120", app.preview_zoom_percent == 120,
               f"zoom={app.preview_zoom_percent}")
        app._on_ctrl_mousewheel_zoom(_wheel(delta=-120))
        _check("滚轮向下缩小一档", app.preview_zoom_percent == 110,
               f"zoom={app.preview_zoom_percent}")

        # Linux/X11 形态：Button-4 上滚 / Button-5 下滚
        app._set_preview_zoom(100)
        app._on_ctrl_mousewheel_zoom(_wheel(delta=0, num=4))
        _check("Button-4 放大", app.preview_zoom_percent == 110,
               f"zoom={app.preview_zoom_percent}")
        app._on_ctrl_mousewheel_zoom(_wheel(delta=0, num=5))
        _check("Button-5 缩小", app.preview_zoom_percent == 100,
               f"zoom={app.preview_zoom_percent}")

        # 高精度鼠标的小幅 delta 也应识别方向
        app._set_preview_zoom(100)
        app._on_ctrl_mousewheel_zoom(_wheel(delta=40))
        _check("小幅正向 delta 放大", app.preview_zoom_percent == 110)
        app._on_ctrl_mousewheel_zoom(_wheel(delta=-40))
        _check("小幅负向 delta 缩小", app.preview_zoom_percent == 100)

        # 缩放标签同步
        _check("百分比标签同步", app.zoom_label.cget("text") == "100%",
               f"text={app.zoom_label.cget('text')!r}")
    finally:
        app.destroy()


def test_ctrl_wheel_zoom_clamp() -> None:
    print("[2] 缩放钳制 50%–200%")
    app = FormatApp()
    app.withdraw()
    try:
        app._set_preview_zoom(195)
        for _ in range(10):
            app._on_ctrl_mousewheel_zoom(_wheel(delta=120))
        _check("放大上限 200%", app.preview_zoom_percent == 200,
               f"zoom={app.preview_zoom_percent}")
        for _ in range(30):
            app._on_ctrl_mousewheel_zoom(_wheel(delta=-120))
        _check("缩小下限 50%", app.preview_zoom_percent == 50,
               f"zoom={app.preview_zoom_percent}")
    finally:
        app.destroy()


def test_plain_wheel_not_hijacked_by_ctrl() -> None:
    print("[3] 按住 Ctrl 时普通滚轮不再翻页")
    app = FormatApp()
    app.withdraw()
    try:
        app.current_view_page = 1
        ret = app._on_preview_mousewheel(_wheel(delta=-120))
        _check("返回 None 以放行更具体的绑定", ret is None, f"ret={ret!r}")
        _check("页面未翻页", app.current_view_page == 1,
               f"page={app.current_view_page}")

        # 不带 Ctrl 的普通滚轮仍走原有翻页/滚动逻辑
        ret2 = app._on_preview_mousewheel(_wheel(delta=-120, state=0))
        _check("无修饰键滚轮仍返回 break", ret2 == "break", f"ret={ret2!r}")
    finally:
        app.destroy()


def test_ctrl_bindings_present() -> None:
    print("[4] 预览区控件均已绑定 Control 快捷键")
    app = FormatApp()
    app.withdraw()
    try:
        widgets = [app.preview_canvas, app.left_class_text]
        widgets.extend(app.preview_text_widgets)
        for w in widgets:
            path = str(w)
            mw = w.bind("<Control-MouseWheel>")
            b4 = w.bind("<Control-Button-4>")
            b5 = w.bind("<Control-Button-5>")
            _check(f"{path} 绑定 Control-MouseWheel", bool(mw), f"binding={mw!r}")
            _check(f"{path} 绑定 Control-Button-4", bool(b4))
            _check(f"{path} 绑定 Control-Button-5", bool(b5))
        # 快捷键提示包含 Ctrl+滚轮
        _check("顶栏提示包含 Ctrl+滚轮",
               "Ctrl+滚轮" in app.shortcut_hint_label.cget("text"))
    finally:
        app.destroy()


def test_zoom_effective_in_fit_width_mode() -> None:
    print("[5] 撑满宽模式下缩放真实生效")
    app = FormatApp()
    app.withdraw()
    try:
        app.on_preview_view_mode_change("双页 · 撑满宽")
        app._set_preview_zoom(100)
        app._draw_a4_page(900, 600)
        w100, _ = app._compute_page_geometry(900, 600)
        app._set_preview_zoom(150)
        w150, _ = app._compute_page_geometry(900, 600)
        _check("150% 页面宽度约为 100% 的 1.5 倍",
               abs(w150 / w100 - 1.5) < 0.05, f"{w100} -> {w150}")

        # 放大到 200% 后页面超出画布宽度，横向滚动条应出现。
        # 注意：withdraw 后 winfo_ismapped 恒为 0，用 winfo_manager 判断 grid 状态。
        app._set_preview_zoom(200)
        app._draw_a4_page(900, 600)
        _check("超宽时横向滚动条显示", app.preview_hscroll.winfo_manager() == "grid",
               f"manager={app.preview_hscroll.winfo_manager()!r}")

        # 回到 100% 后滚动条收起
        app._set_preview_zoom(100)
        app._draw_a4_page(900, 600)
        _check("未超宽时横向滚动条收起",
               app.preview_hscroll.winfo_manager() == "",
               f"manager={app.preview_hscroll.winfo_manager()!r}")
    finally:
        app.destroy()


def test_event_generate_end_to_end() -> None:
    print("[6] 端到端：event_generate 触发缩放")
    with tempfile.TemporaryDirectory() as tmp:
        f = Path(tmp) / "doc.docx"
        doc = Document()
        for p in range(3):
            doc.add_paragraph(f"第{p}章")
            doc.add_paragraph("内容" * 30)
        doc.save(str(f))
        app = FormatApp()
        app.withdraw()
        try:
            app.docx_handler.load(str(f))
            app._refresh_document_view()
            app._set_preview_zoom(100)
            app.preview_canvas.event_generate(
                "<Control-MouseWheel>", x=10, y=10, delta=120
            )
            app.update_idletasks()
            _check("画布上 Ctrl+滚轮事件放大", app.preview_zoom_percent == 110,
                   f"zoom={app.preview_zoom_percent}")
            app.preview_text_widgets[0].event_generate(
                "<Control-MouseWheel>", x=5, y=5, delta=-120
            )
            app.update_idletasks()
            _check("页面文本上 Ctrl+滚轮事件缩小", app.preview_zoom_percent == 100,
                   f"zoom={app.preview_zoom_percent}")
        finally:
            app.destroy()


if __name__ == "__main__":
    print("=== 批次6：Ctrl+滚轮缩放 ===")
    test_ctrl_wheel_zoom_steps()
    test_ctrl_wheel_zoom_clamp()
    test_plain_wheel_not_hijacked_by_ctrl()
    test_ctrl_bindings_present()
    test_zoom_effective_in_fit_width_mode()
    test_event_generate_end_to_end()
    print(f"\n批次6结果: {PASS} 通过, {FAIL} 失败")
    sys.exit(1 if FAIL else 0)
