"""批次 1 回归测试：数据损坏类修复。

覆盖 4 个修复点：
1. load()/save() 不再把全文行距强制改写为 1.25（仅打开文档即损坏的问题）；
2. 工具栏"固定18pt"真正应用 EXACTLY 18pt，且不被分页规则改写；
3. 参考文献格式化不再一律补 [J]（按内容推断类型，判断不出则保持原样）；
4. 目录重复更新不再累积折页空段（空白页）。

运行方式（项目根目录）：
    python tests/test_batch1_data_integrity.py
"""

from __future__ import annotations

import re
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt

from src.docx_handler import DocxHandler


def _make_spacing_doc(path: Path) -> None:
    doc = Document()
    p1 = doc.add_paragraph("单倍行距段落")
    p1.paragraph_format.line_spacing = 1.0
    p2 = doc.add_paragraph("双倍行距段落")
    p2.paragraph_format.line_spacing = 2.0
    p3 = doc.add_paragraph("固定值20pt段落")
    p3.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p3.paragraph_format.line_spacing = Pt(20)
    p4 = doc.add_paragraph("1.5倍行距段落")
    p4.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p4.paragraph_format.line_spacing = 1.5
    doc.save(str(path))


def _make_toc_doc(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("摘要")
    doc.add_paragraph("本文研究了某问题。")
    doc.add_paragraph("关键词：测试；回归")
    doc.add_paragraph("第 1 章 绪论")
    doc.add_paragraph("第一章内容。")
    doc.add_paragraph("1.1 小节标题")
    doc.add_paragraph("小节内容。")
    doc.add_paragraph("第 2 章 结论")
    doc.add_paragraph("第二章内容。")
    doc.save(str(path))


def _make_reference_doc(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("第 1 章 绪论")
    doc.add_paragraph("第一章内容。")
    doc.add_paragraph("参考文献")
    doc.add_paragraph("[1] 张三. 教育学原理. 北京: 人民教育出版社, 2019.")
    doc.add_paragraph("[2] 王五. 某高校人才培养研究[学位论文]. 北京: 某大学, 2021.")
    doc.add_paragraph("[3] 李四. 某学术会议论文集. 北京: 某出版社, 2020.")
    doc.add_paragraph("[4] 赵六. 某领域研究综述. 计算机学报, 2020, 43(3): 45-60.")
    doc.add_paragraph("[5] 孙七. 一条没有任何类型线索的条目. 2020.")
    doc.add_paragraph("[6] Smith J. Title of a book. MIT Press, 2018.")
    doc.add_paragraph("致谢")
    doc.add_paragraph("感谢。")
    doc.save(str(path))


def _count_break_only_paragraphs(doc: Document) -> int:
    """统计无文字、仅含分页符的段落（目录折页空段）。"""
    count = 0
    for para in doc.paragraphs:
        if (para.text or "").strip():
            continue
        has_break = any(
            node.get(qn("w:type")) == "page" for node in para._p.xpath(".//w:br")
        )
        if has_break:
            count += 1
    return count


def test_load_does_not_rewrite_line_spacing() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        src = tmp_dir / "spacing.docx"
        _make_spacing_doc(src)

        handler = DocxHandler()
        handler.load(str(src))
        fmts = [p.paragraph_format for p in handler._document.paragraphs]
        assert fmts[0].line_spacing == 1.0, f"单倍行距被改写: {fmts[0].line_spacing}"
        assert fmts[1].line_spacing == 2.0, f"双倍行距被改写: {fmts[1].line_spacing}"
        assert fmts[2].line_spacing_rule == WD_LINE_SPACING.EXACTLY, "固定值行距规则被改写"
        assert fmts[2].line_spacing == Pt(20), f"固定值行距被改写: {fmts[2].line_spacing}"
        assert fmts[3].line_spacing == 1.5, f"1.5倍行距被改写: {fmts[3].line_spacing}"

        # 保存后重新读取，格式必须保持不变。
        out = tmp_dir / "spacing_saved.docx"
        handler.save(str(out))
        reloaded = Document(str(out))
        rfmts = [p.paragraph_format for p in reloaded.paragraphs]
        assert rfmts[0].line_spacing == 1.0, "保存后单倍行距被改写"
        assert rfmts[1].line_spacing == 2.0, "保存后双倍行距被改写"
        assert rfmts[2].line_spacing_rule == WD_LINE_SPACING.EXACTLY, "保存后固定值规则被改写"
        assert rfmts[2].line_spacing == Pt(20), "保存后固定值行距被改写"
        assert rfmts[3].line_spacing == 1.5, "保存后1.5倍行距被改写"
    print("PASS test_load_does_not_rewrite_line_spacing")


def test_fixed18pt_is_exact_18pt() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        src = tmp_dir / "fixed18.docx"
        doc = Document()
        doc.add_paragraph("测试段落")
        doc.save(str(src))

        handler = DocxHandler()
        handler.load(str(src))
        changed = handler.apply_line_spacing({0}, "fixed18pt")
        assert changed == 1, "固定18pt 应用失败"
        fmt = handler._document.paragraphs[0].paragraph_format
        assert fmt.line_spacing_rule == WD_LINE_SPACING.EXACTLY, (
            f"固定18pt 未使用 EXACTLY: {fmt.line_spacing_rule}"
        )
        assert fmt.line_spacing == Pt(18), f"固定18pt 值错误: {fmt.line_spacing}"

        # 分页规则不得改写已应用的固定行距。
        handler.enforce_word_pagination_rules()
        fmt = handler._document.paragraphs[0].paragraph_format
        assert fmt.line_spacing_rule == WD_LINE_SPACING.EXACTLY, "分页规则改写了固定行距规则"
        assert fmt.line_spacing == Pt(18), "分页规则改写了固定行距值"

        # 保存后重新读取，固定行距必须保留。
        out = tmp_dir / "fixed18_saved.docx"
        handler.save(str(out))
        reloaded = Document(str(out))
        rfmt = reloaded.paragraphs[0].paragraph_format
        assert rfmt.line_spacing_rule == WD_LINE_SPACING.EXACTLY, "保存后固定行距规则丢失"
        assert rfmt.line_spacing == Pt(18), f"保存后固定行距值错误: {rfmt.line_spacing}"
    print("PASS test_fixed18pt_is_exact_18pt")


def test_reference_type_inference() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "refs.docx"
        _make_reference_doc(src)
        handler = DocxHandler()
        handler.load(str(src))
        result = handler.format_references_gbt(mode="order")
        assert result.get("ok"), f"参考文献格式化失败: {result}"
        texts = [p.text for p in handler._document.paragraphs]
        joined = "\n".join(texts)

        assert any("人民教育出版社" in t and "[M]" in t for t in texts), (
            f"专著未推断为 [M]: {[t for t in texts if '教育出版社' in t]}"
        )
        assert any("学位论文" in t and "[D]" in t for t in texts), (
            f"学位论文未推断为 [D]: {[t for t in texts if '学位论文' in t]}"
        )
        assert any("学术会议论文集" in t and "[C]" in t for t in texts), (
            f"会议论文集未推断为 [C]: {[t for t in texts if '学术会议' in t]}"
        )
        assert any("计算机学报" in t and "[J]" in t for t in texts), (
            f"期刊未推断为 [J]: {[t for t in texts if '学报' in t]}"
        )
        assert any("MIT Press" in t and "[M]" in t for t in texts), (
            f"英文专著未推断为 [M]: {[t for t in texts if 'MIT Press' in t]}"
        )
        # 无任何类型线索的条目：不得补任何标识，更不得补 [J]。
        vague = [t for t in texts if "没有任何类型线索" in t]
        assert vague, "未找到无标识线索的测试条目"
        assert "[J]" not in vague[0], f"无线索条目被误注入 [J]: {vague[0]}"
        assert not any(m in vague[0] for m in ("[M]", "[D]", "[C]", "[R]", "[S]", "[P]")), (
            f"无线索条目被误注入类型标识: {vague[0]}"
        )
        # 全文不得出现“句末标点后又追加类型标识”的污染形态（如 “。 [J]”）。
        assert not re.search(r"[。.]\s*\[[A-Z]", joined), (
            f"存在句末追加类型标识的污染: {joined!r}"
        )
        # 合法形态：标识紧贴句末句号之前，如 “45-60 [J].”
        assert any(re.search(r"\d\s\[J\]\.$", t) for t in texts), (
            f"期刊条目未按 “…页码 [J].” 形态输出: {[t for t in texts if '学报' in t]}"
        )
    print("PASS test_reference_type_inference")


def test_toc_update_does_not_accumulate_breaks() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "toc.docx"
        _make_toc_doc(src)
        handler = DocxHandler()
        handler.load(str(src))

        first = handler.auto_generate_toc(levels=3)
        assert first.get("ok"), f"目录生成失败: {first}"
        doc = handler._document
        breaks_after_first = _count_break_only_paragraphs(doc)
        assert breaks_after_first == 2, (
            f"首次生成后期望 2 个折页空段（目录前/后各1），实际 {breaks_after_first}"
        )
        paras_after_first = len(doc.paragraphs)
        entries_after_first = sum(1 for p in doc.paragraphs if p.text.startswith("第 1 章 绪论\t"))
        assert entries_after_first == 1, f"目录条目重复: {entries_after_first}"

        for round_no in (2, 3, 4):
            upd = handler.update_toc()
            assert upd.get("ok"), f"第 {round_no} 次更新目录失败: {upd}"
            breaks = _count_break_only_paragraphs(handler._document)
            assert breaks == 2, f"第 {round_no} 次更新后折页空段累积: {breaks}"
            paras = len(handler._document.paragraphs)
            assert paras == paras_after_first, (
                f"第 {round_no} 次更新后段落数漂移: {paras} != {paras_after_first}"
            )
            entries = sum(
                1 for p in handler._document.paragraphs if p.text.startswith("第 1 章 绪论\t")
            )
            assert entries == 1, f"第 {round_no} 次更新后目录条目重复: {entries}"
    print("PASS test_toc_update_does_not_accumulate_breaks")


def main() -> int:
    tests = [
        test_load_does_not_rewrite_line_spacing,
        test_fixed18pt_is_exact_18pt,
        test_reference_type_inference,
        test_toc_update_does_not_accumulate_breaks,
    ]
    failed = 0
    for test in tests:
        try:
            test()
        except AssertionError as exc:
            failed += 1
            print(f"FAIL {test.__name__}: {exc}")
        except Exception as exc:  # noqa: BLE001
            failed += 1
            print(f"ERROR {test.__name__}: {type(exc).__name__}: {exc}")
    print(f"\n批次1回归：{len(tests) - failed}/{len(tests)} 通过")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
