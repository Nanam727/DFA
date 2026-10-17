"""批次 2 回归测试：显示与交互类修复。

覆盖 4 个修复点：
1. 状态栏真实存在，_update_hint 的反馈立即可见（不再静默丢弃）；
2. 空页占位符不再堆叠（整对空页只显示一条；空文档右页只显示一遍）；
3. “每页X行/每行Y字”网格信息标签真实存在且随边距更新；
4. Ctrl+S 等快捷键改绑 ui 自身实现，controller 双实现已移除，保存不再崩溃。

运行方式（项目根目录）：
    python tests/test_batch2_display.py
"""

from __future__ import annotations

import re
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from docx.enum.text import WD_BREAK

import src.ui as ui_module
from src.ui import FormatApp


def _make_gap_doc(path: Path, breaks: int = 4) -> None:
    """第 1 页有内容；一个段落内多个分页符使页码跳空（breaks=4 时 2/3/4 页全空）。"""
    doc = Document()
    doc.add_paragraph("第 1 章 绪论")
    doc.add_paragraph("第一章内容。")
    p = doc.add_paragraph("")
    for _ in range(breaks):
        p.add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph("第 2 章 结论")
    doc.add_paragraph("第二章内容。")
    doc.save(str(path))


def _widget_text(widget) -> str:
    return widget.get("1.0", "end")


def test_status_bar_visible() -> None:
    app = FormatApp()
    app.withdraw()
    try:
        assert hasattr(app, "status_text_label"), "状态栏控件未创建"
        app._update_hint("已应用字号 12pt")
        label_text = app.status_text_label.cget("text")
        assert "已应用字号 12pt" in label_text, f"状态栏未显示反馈: {label_text!r}"
        assert app.status_text == label_text, "status_text 属性与控件不同步"
        app._update_hint("已全选")
        assert "已全选" in app.status_text_label.cget("text"), "状态栏未更新第二条反馈"
    finally:
        app.destroy()
    print("PASS test_status_bar_visible")


def test_placeholder_no_stacking() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        empty = Path(tmp) / "empty.docx"
        Document().save(str(empty))

        app = FormatApp()
        app.withdraw()
        try:
            # 场景 A：整对空页（breaks=4 -> 页码 1 跳到 5，视图页 (3,4) 全空）。
            gap = Path(tmp) / "gap4.docx"
            _make_gap_doc(gap, breaks=4)
            app.docx_handler.load(str(gap))
            app._refresh_document_view()
            app.current_view_page = 3
            app._refresh_document_view(preserve_selection=True)
            left_text = _widget_text(app.preview_text_widgets[0])
            right_text = _widget_text(app.preview_text_widgets[1])
            assert "（当前页无可显示段落）" in left_text, f"空页缺少提示: {left_text!r}"
            assert "（左页无内容）" not in left_text, f"占位符堆叠: {left_text!r}"
            assert right_text.count("（右页无内容）") == 1, f"右页占位异常: {right_text!r}"

            # 场景 A2：仅左页为空（breaks=2 -> 视图页 (2,3)，第 3 页有内容）。
            gap2 = Path(tmp) / "gap2.docx"
            _make_gap_doc(gap2, breaks=2)
            app.docx_handler.load(str(gap2))
            app._refresh_document_view()
            app.current_view_page = 2
            app._refresh_document_view(preserve_selection=True)
            left_text = _widget_text(app.preview_text_widgets[0])
            assert left_text.count("（左页无内容）") == 1, f"左页空页提示异常: {left_text!r}"
            assert "（当前页无可显示段落）" not in left_text, f"占位符堆叠: {left_text!r}"
            assert "第 2 章 结论" in _widget_text(app.preview_text_widgets[1]), (
                "右页内容缺失"
            )

            # 场景 B：空文档，右页提示只出现一次。
            app.docx_handler.load(str(empty))
            app._refresh_document_view()
            left_text = _widget_text(app.preview_text_widgets[0])
            right_text = _widget_text(app.preview_text_widgets[1])
            assert "（文档无可显示段落）" in left_text, f"空文档左页提示缺失: {left_text!r}"
            assert right_text.count("（右页无内容）") == 1, (
                f"空文档右页提示重复: {right_text!r}"
            )
        finally:
            app.destroy()
    print("PASS test_placeholder_no_stacking")


def test_grid_info_label() -> None:
    app = FormatApp()
    app.withdraw()
    try:
        assert hasattr(app, "grid_info_label"), "网格信息标签未创建"
        text = app.grid_info_label.cget("text")
        assert re.search(r"每页 \d+ 行 / 每行 \d+ 字", text), f"网格信息格式异常: {text!r}"
        # 切换边距预设后应刷新。
        app.current_margin_preset_name = "窄"
        app._update_grid_info(app._get_active_margins_mm())
        text = app.grid_info_label.cget("text")
        assert "窄" in text and re.search(r"每页 \d+ 行 / 每行 \d+ 字", text), (
            f"切换边距后网格信息未更新: {text!r}"
        )
    finally:
        app.destroy()
    print("PASS test_grid_info_label")


def test_save_shortcut_works_without_controller() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "doc.docx"
        doc = Document()
        doc.add_paragraph("测试段落")
        doc.save(str(src))

        app = FormatApp()
        app.withdraw()
        try:
            assert not hasattr(app, "controller"), "controller 双实现未移除"
            app.docx_handler.load(str(src))
            before = src.stat().st_mtime_ns

            # 屏蔽模态框，直接验证 Ctrl+S 链路。
            original_showinfo = ui_module.messagebox.showinfo
            ui_module.messagebox.showinfo = lambda *a, **k: None
            try:
                result = app.on_save_shortcut(None)
            finally:
                ui_module.messagebox.showinfo = original_showinfo

            assert result == "break", f"快捷键返回值异常: {result!r}"
            assert src.stat().st_mtime_ns != before, "Ctrl+S 未实际保存文档"
        finally:
            app.destroy()
    print("PASS test_save_shortcut_works_without_controller")


def main() -> int:
    tests = [
        test_status_bar_visible,
        test_placeholder_no_stacking,
        test_grid_info_label,
        test_save_shortcut_works_without_controller,
    ]
    failed = 0
    for test in tests:
        try:
            test()
        except AssertionError as exc:
            failed += 1
            print(f"FAIL {test.__name__}: {exc}")
        except Exception as exc:  # noqa: BLE001
            failed += 1
            print(f"ERROR {test.__name__}: {type(exc).__name__}: {exc}")
    print(f"\n批次2回归：{len(tests) - failed}/{len(tests)} 通过")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
