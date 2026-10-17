"""批次5回归：浏览区优化（工具栏折叠 / 浏览模式 / 缩放 / 滚动翻页）。

运行：python tests/test_batch5_preview_ux.py
"""

from __future__ import annotations

import sys
import tempfile
import types
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from docx.enum.text import WD_BREAK

from src.ui import FormatApp

PASS = "PASS"
FAIL = "FAIL"
_results: list[tuple[str, str, str]] = []


def _record(name: str, ok: bool, detail: str = "") -> None:
    _results.append((name, PASS if ok else FAIL, detail))
    print(f"{PASS if ok else FAIL} {name}" + (f" :: {detail}" if detail else ""))


def _check(name: str, condition: bool, detail: str = "") -> bool:
    _record(name, bool(condition), detail)
    return bool(condition)


def _wheel(delta: int):
    """构造一个滚轮事件。"""
    return types.SimpleNamespace(delta=delta, num=0)


def _make_multipage_doc(path: Path, pages: int = 6) -> None:
    doc = Document()
    for page in range(1, pages + 1):
        doc.add_paragraph(f"第 {page} 章 标题{page}")
        doc.add_paragraph(f"第 {page} 章正文内容。" * 6)
        if page < pages:
            para = doc.add_paragraph("")
            para.add_run().add_break(WD_BREAK.PAGE)
    doc.save(str(path))


def test_toggle_right_panel() -> None:
    app = FormatApp()
    app.withdraw()
    try:
        _check("初始工具栏展开", app._right_panel_visible is True)
        _check(
            "初始按钮文案为收起",
            app.btn_toggle_tools.cget("text") == "收起",
        )
        _check(
            "初始右栏由 grid 管理",
            app.right_panel_frame.winfo_manager() == "grid",
        )

        app.on_toggle_right_panel()
        _check("收起后按钮文案为展开", app.btn_toggle_tools.cget("text") == "展开")
        _check(
            "收起后右栏脱离 grid",
            app.right_panel_frame.winfo_manager() == "",
            f"manager={app.right_panel_frame.winfo_manager()!r}",
        )
        _check("收起后状态标记为 False", app._right_panel_visible is False)

        app.on_toggle_right_panel()
        _check("再次点击恢复展开", app._right_panel_visible is True)
        _check(
            "恢复后右栏由 grid 管理",
            app.right_panel_frame.winfo_manager() == "grid",
        )
        _check("恢复后按钮文案为收起", app.btn_toggle_tools.cget("text") == "收起")
    finally:
        app.destroy()


def test_page_geometry_modes() -> None:
    app = FormatApp()
    app.withdraw()
    try:
        # 双页整页：受高度上限约束
        app.preview_view_mode = "pair_fit"
        fit_w, fit_h = app._compute_page_geometry(900, 600)
        _check("整页模式为双页", app._preview_pages() == 2)
        _check(
            "整页模式页面受高度约束",
            fit_w == 410 and fit_h == 579,
            f"page={fit_w}x{fit_h}",
        )

        # 双页撑满宽：每页撑满半栏
        app.preview_view_mode = "pair_width"
        pw_w, pw_h = app._compute_page_geometry(900, 600)
        _check("双页撑满宽仍为双页", app._preview_pages() == 2)
        _check(
            "双页撑满宽页面更大",
            pw_w == 435 and pw_h == 615,
            f"page={pw_w}x{pw_h}",
        )

        # 单页撑满宽：整页撑满，文字最大
        app.preview_view_mode = "single_width"
        sw_w, sw_h = app._compute_page_geometry(900, 600)
        _check("单页模式页数为 1", app._preview_pages() == 1)
        _check(
            "单页撑满宽页面最大",
            sw_w == 880 and sw_h == 1244,
            f"page={sw_w}x{sw_h}",
        )
        _check(
            "三种模式页面尺寸递增",
            fit_w < pw_w < sw_w,
            f"{fit_w} < {pw_w} < {sw_w}",
        )
    finally:
        app.destroy()


def test_zoom_controls() -> None:
    app = FormatApp()
    app.withdraw()
    try:
        app.preview_view_mode = "pair_fit"
        app._sync_zoom_controls_state()
        _check(
            "整页模式缩放按钮可用",
            app.btn_zoom_in.cget("state") == "normal"
            and app.btn_zoom_out.cget("state") == "normal",
        )
        _check("初始缩放 100%", app.zoom_label.cget("text") == "100%")

        app.on_zoom_in()
        _check("放大到 110%", app.preview_zoom_percent == 110)
        _check("标签同步", app.zoom_label.cget("text") == "110%")

        for _ in range(20):
            app.on_zoom_out()
        _check("缩小下限 50%", app.preview_zoom_percent == 50)
        for _ in range(30):
            app.on_zoom_in()
        _check("放大上限 200%", app.preview_zoom_percent == 200)

        # 整页模式：页面必须完整可见，放大受高度上限约束（1000x500 恒为 339）
        app.preview_zoom_percent = 100
        base_w, _ = app._compute_page_geometry(1000, 500)
        app._set_preview_zoom(140)
        capped_w, _ = app._compute_page_geometry(1000, 500)
        _check(
            "整页模式放大受高度上限约束",
            base_w == 339 and capped_w == 339,
            f"{base_w} -> {capped_w}",
        )

        # 撑满宽模式：放大真正放大页面，超出部分横向滚动查看
        app.preview_view_mode = "pair_width"
        app.preview_zoom_percent = 100
        w100, _ = app._compute_page_geometry(900, 600)
        app._set_preview_zoom(200)
        w200, _ = app._compute_page_geometry(900, 600)
        _check(
            "撑满宽模式放大页面翻倍",
            w100 == 435 and w200 == 870,
            f"{w100} -> {w200}",
        )
        _check(
            "撑满宽模式缩放按钮可用",
            app.btn_zoom_in.cget("state") == "normal"
            and app.btn_zoom_out.cget("state") == "normal",
        )
        _check("撑满宽模式标签显示百分比", app.zoom_label.cget("text") == "200%")
    finally:
        app.destroy()


def test_wheel_scroll_and_page_turn() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        doc_path = Path(tmp) / "multi.docx"
        _make_multipage_doc(doc_path, pages=6)

        app = FormatApp()
        app.withdraw()
        try:
            app.docx_handler.load(str(doc_path))
            app._refresh_document_view()
            _check("文档共 6 页", app.total_view_pages == 6, f"total={app.total_view_pages}")

            # 双页撑满宽：页面高于画布 -> 滚轮滚动，滚到底再滚才翻页
            app.preview_view_mode = "pair_width"
            app._draw_a4_page(900, 600)
            _check("撑满宽模式可滚动", app._preview_scrollable() is True)
            _check(
                "撑满宽模式显示滚动条",
                app.preview_vscroll.winfo_manager() == "grid",
            )

            app.preview_canvas.yview_moveto(0.0)
            app._on_preview_mousewheel(_wheel(-1))
            scrolled = app.preview_canvas.yview()[0]
            _check("滚轮向下为滚动而非翻页", scrolled > 0 and app.current_view_page == 1,
                   f"yview={scrolled}, page={app.current_view_page}")

            app.preview_canvas.yview_moveto(1.0)
            app._on_preview_mousewheel(_wheel(-1))
            _check(
                "滚到底再向下翻到下一页对",
                app.current_view_page == 3,
                f"page={app.current_view_page}",
            )
            _check(
                "翻页后回到页顶",
                app.preview_canvas.yview()[0] <= 0.001,
            )

            # 双页整页：不可滚动，滚轮直接翻页（保持原有行为）
            app.preview_view_mode = "pair_fit"
            app._draw_a4_page(900, 600)
            _check("整页模式不显示滚动条",
                   app.preview_vscroll.winfo_manager() == "",
                   f"manager={app.preview_vscroll.winfo_manager()!r}")
            _check("整页模式不可滚动", app._preview_scrollable() is False)
            app._on_preview_mousewheel(_wheel(-1))
            _check("整页模式滚轮翻页 3->5", app.current_view_page == 5,
                   f"page={app.current_view_page}")
            app._on_preview_mousewheel(_wheel(1))
            _check("整页模式滚轮回翻 5->3", app.current_view_page == 3,
                   f"page={app.current_view_page}")
        finally:
            app.destroy()


def test_single_mode_layout() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        doc_path = Path(tmp) / "multi.docx"
        _make_multipage_doc(doc_path, pages=4)

        app = FormatApp()
        app.withdraw()
        try:
            app.docx_handler.load(str(doc_path))
            app._refresh_document_view()

            # 先在双页模式下创建右页窗口
            app.preview_view_mode = "pair_width"
            app._draw_a4_page(900, 600)
            win_right = app.page_window_ids[1]
            _check("双页模式右页窗口可见",
                   app.preview_canvas.itemcget(win_right, "state") == "normal")

            # 切到单页撑满宽
            app.on_preview_view_mode_change("单页 · 撑满宽")
            app._draw_a4_page(900, 600)
            _check("单页模式右页窗口隐藏",
                   app.preview_canvas.itemcget(win_right, "state") == "hidden",
                   f"state={app.preview_canvas.itemcget(win_right, 'state')!r}")
            _check("单页页码指示器", app.page_info_label.cget("text") == "第 1 / 4 页",
                   f"text={app.page_info_label.cget('text')!r}")

            # 单页模式滚轮翻页步进为 1
            app.preview_canvas.yview_moveto(1.0)
            app._on_preview_mousewheel(_wheel(-1))
            _check("单页模式翻页步进 1", app.current_view_page == 2,
                   f"page={app.current_view_page}")
            _check("单页页码指示器更新",
                   app.page_info_label.cget("text") == "第 2 / 4 页")

            # 单页文字显著大于双页（可读性代理指标，同一画布尺寸下比较）
            single_scale = app._preview_font_scale_value
            app.on_preview_view_mode_change("双页 · 整页")
            app._draw_a4_page(900, 600)
            _check("切回双页右页恢复可见",
                   app.preview_canvas.itemcget(win_right, "state") == "normal")
            pair_scale = app._preview_font_scale_value
            _check(
                "单页字号缩放大于双页",
                single_scale > pair_scale,
                f"single={single_scale:.3f} > pair={pair_scale:.3f}",
            )
        finally:
            app.destroy()


def test_view_mode_change_updates_state() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        doc_path = Path(tmp) / "multi.docx"
        _make_multipage_doc(doc_path, pages=6)
        app = FormatApp()
        app.withdraw()
        try:
            app.docx_handler.load(str(doc_path))
            app._refresh_document_view()
            app.preview_view_mode = "pair_fit"
            app.on_preview_view_mode_change("双页 · 撑满宽")
            _check("模式状态已更新", app.preview_view_mode == "pair_width")
            _check("下拉显示同步",
                   app.view_mode_combo.get() == "双页 · 撑满宽")
            # 无效选择被忽略
            app.on_preview_view_mode_change("不存在的模式")
            _check("无效模式被忽略", app.preview_view_mode == "pair_width")
            # 单页模式下偶数页对齐到奇数起始页
            app.current_view_page = 4
            app.on_preview_view_mode_change("单页 · 撑满宽")
            _check("单页模式偶数页前移一页", app.current_view_page == 3,
                   f"page={app.current_view_page}")
        finally:
            app.destroy()


def main() -> int:
    test_toggle_right_panel()
    test_page_geometry_modes()
    test_zoom_controls()
    test_wheel_scroll_and_page_turn()
    test_single_mode_layout()
    test_view_mode_change_updates_state()
    failed = [r for r in _results if r[1] == FAIL]
    print(f"\n批次5回归：{len(_results) - len(failed)}/{len(_results)} 通过")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
