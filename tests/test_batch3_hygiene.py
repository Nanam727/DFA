"""批次 3 回归测试：代码卫生与中小问题修复。

覆盖：
1. 对齐无变化时不压入撤销快照、不计成功；
2. 反向查找全新搜索时从最后 1 处起跳（不再落到倒数第 2 处）；
3. 清理操作仍可经主撤销栈撤销（清理快照死机制移除后）；
4. 直接保存失败时先说明原因再引导另存为（不再静默弹窗）；
5. 右侧面板所有按钮的 command 处理函数真实存在（死代码移除后的接线冒烟）。

运行方式（项目根目录）：
    python tests/test_batch3_hygiene.py
"""

from __future__ import annotations

import re
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

import src.ui as ui_module
from src.docx_handler import DocxHandler
from src.ui import FormatApp


def test_alignment_no_change_no_undo_push() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "align.docx"
        doc = Document()
        p = doc.add_paragraph("段落")
        p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
        doc.save(str(src))

        handler = DocxHandler()
        handler.load(str(src))
        before = handler.get_undo_redo_counts()["undo"]

        assert handler.apply_alignment({0}, "left") == 0, "无变化时仍报告成功"
        assert handler.get_undo_redo_counts()["undo"] == before, "无变化时仍压入撤销快照"

        assert handler.apply_alignment({0}, "center") == 1, "实际变化未应用"
        assert handler.get_undo_redo_counts()["undo"] == before + 1, "实际变化未记录撤销"
        assert handler._document.paragraphs[0].alignment == WD_PARAGRAPH_ALIGNMENT.CENTER
    print("PASS test_alignment_no_change_no_undo_push")


def test_reverse_find_first_jump() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "find.docx"
        doc = Document()
        for i in range(4):
            doc.add_paragraph(f"第 {i + 1} 段包含关键词测试")
        doc.save(str(src))

        app = FormatApp()
        app.withdraw()
        try:
            app.docx_handler.load(str(src))
            app._refresh_document_view()

            # 反向查找：全新搜索应从最后 1 处起跳。
            app._find_and_focus_next("关键词", direction=-1)
            total = len(app.find_match_indices)
            assert total == 4, f"匹配数异常: {total}"
            assert app.find_match_cursor == total - 1, (
                f"反向查找首次跳错位置: cursor={app.find_match_cursor}, 期望 {total - 1}"
            )

            # 反向续查：循环到上一个匹配（3 -> 2）。
            app._find_and_focus_next("关键词", direction=-1)
            assert app.find_match_cursor == 2, (
                f"反向续查位置错误: cursor={app.find_match_cursor}, 期望 2"
            )

            # 正向查找：全新搜索应从第 1 处起跳。
            app.find_last_signature = None  # 强制重新搜索
            app._find_and_focus_next("关键词", direction=1)
            assert app.find_match_cursor == 0, (
                f"正向查找首次跳错位置: cursor={app.find_match_cursor}, 期望 0"
            )

            # 正向续查：1 -> 2。
            app._find_and_focus_next("关键词", direction=1)
            assert app.find_match_cursor == 1, (
                f"正向续查位置错误: cursor={app.find_match_cursor}, 期望 1"
            )
        finally:
            app.destroy()
    print("PASS test_reverse_find_first_jump")


def test_clean_remove_empty_still_undoable() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "clean.docx"
        doc = Document()
        doc.add_paragraph("内容段")
        doc.add_paragraph("")
        doc.add_paragraph("")
        doc.add_paragraph("内容段二")
        doc.save(str(src))

        handler = DocxHandler()
        handler.load(str(src))
        assert handler.count_empty_paragraphs() == 2
        assert handler.clean_remove_empty_paragraphs() == 2
        assert handler.count_empty_paragraphs() == 0, "空段落未删除"

        assert handler.undo_last_action(), "清理操作无法撤销"
        assert handler.count_empty_paragraphs() == 2, "撤销后空段落未恢复"
    print("PASS test_clean_remove_empty_still_undoable")


def test_save_failure_shows_reason() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "save.docx"
        doc = Document()
        doc.add_paragraph("测试")
        doc.save(str(src))

        app = FormatApp()
        app.withdraw()
        try:
            app.docx_handler.load(str(src))

            def _boom(*_args, **_kwargs):
                raise RuntimeError("磁盘空间不足")

            app.docx_handler.save = _boom

            errors: list[tuple[str, str]] = []
            original_showerror = ui_module.messagebox.showerror
            ui_module.messagebox.showerror = lambda title, msg, *a, **k: errors.append(
                (title, msg)
            )
            save_as_called: list[bool] = []
            original_save_as = app.on_save_as_document
            app.on_save_as_document = lambda: save_as_called.append(True)
            try:
                app.on_save_document()
            finally:
                ui_module.messagebox.showerror = original_showerror
                app.on_save_as_document = original_save_as

            assert errors, "保存失败时没有给出任何提示"
            _title, message = errors[0]
            assert "磁盘空间不足" in message, f"未说明失败原因: {message!r}"
            assert save_as_called == [True], "失败后未引导另存为"
        finally:
            app.destroy()
    print("PASS test_save_failure_shows_reason")


def test_all_button_handlers_exist() -> None:
    """右侧面板/工具栏所有 command=self.xxx 的处理器必须真实存在。"""
    ui_source = Path("src/ui.py").read_text(encoding="utf-8")
    handler_names = sorted(set(re.findall(r"command=self\.(\w+)", ui_source)))
    assert handler_names, "未抽取到任何按钮处理器"

    app = FormatApp()
    app.withdraw()
    try:
        missing = [name for name in handler_names if not hasattr(app, name)]
        assert not missing, f"按钮绑定了不存在的处理器: {missing}"
    finally:
        app.destroy()
    print(f"PASS test_all_button_handlers_exist（{len(handler_names)} 个处理器）")


def main() -> int:
    tests = [
        test_alignment_no_change_no_undo_push,
        test_reverse_find_first_jump,
        test_clean_remove_empty_still_undoable,
        test_save_failure_shows_reason,
        test_all_button_handlers_exist,
    ]
    failed = 0
    for test in tests:
        try:
            test()
        except AssertionError as exc:
            failed += 1
            print(f"FAIL {test.__name__}: {exc}")
        except Exception as exc:  # noqa: BLE001
            failed += 1
            print(f"ERROR {test.__name__}: {type(exc).__name__}: {exc}")
    print(f"\n批次3回归：{len(tests) - failed}/{len(tests)} 通过")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
