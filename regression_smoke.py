from __future__ import annotations

from pathlib import Path
import tempfile

from docx import Document

from src.docx_handler import DocxHandler


def main() -> None:
    path = Path(tempfile.gettempdir()) / "adfa_regression_smoke.docx"
    doc = Document()
    doc.add_paragraph("摘要")
    doc.add_paragraph("第 1 章 绪论")
    doc.add_paragraph("正文内容 test test")
    doc.add_paragraph("参考文献")
    doc.add_paragraph("[1] 张三. 题名 [J]. 刊名, 2020.")
    doc.add_paragraph("致谢")
    doc.save(path)

    handler = DocxHandler()
    handler.load(path)
    assert handler.find_text("test")
    rep = handler.replace_text("test", "demo", max_replacements_per_para=1)
    assert rep["replaced_count"] >= 1
    assert isinstance(handler.one_click_section_and_page_number(), dict)
    assert isinstance(handler.auto_generate_toc(levels=3, mode="smart"), dict)
    assert isinstance(handler.format_references_gbt("order"), dict)
    assert handler.undo_last_action() in (True, False)
    assert handler.redo_last_action() in (True, False)
    handler.save(path)
    print("REGRESSION_SMOKE_OK")


if __name__ == "__main__":
    main()
